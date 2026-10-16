import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT_DIR = Path(__file__).resolve().parents[2] / "outputs" / "candidate_dataset"
FULL_CSV = OUTPUT_DIR / "talentnet_candidates_500_full.csv"
IMPORT_CSV = OUTPUT_DIR / "talentnet_candidates_500_import.csv"
SUMMARY_JSON = OUTPUT_DIR / "dataset_summary.json"
OUTPUT_XLSX = OUTPUT_DIR / "talentnet_candidates_500_dataset.xlsx"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_sheet(ws, headers: list[str], rows: list[dict[str, str]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    for column_cells in ws.columns:
        header = str(column_cells[0].value or "")
        width = min(max(len(header) + 4, 12), 42)
        ws.column_dimensions[column_cells[0].column_letter].width = width
    table = Table(displayName=ws.title.replace(" ", "") + "Table", ref=ws.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def main() -> None:
    full_rows = read_csv(FULL_CSV)
    import_rows = read_csv(IMPORT_CSV)
    summary = json.loads(SUMMARY_JSON.read_text(encoding="utf-8"))

    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard"
    dashboard["A1"] = "TalentNet Synthetic Candidate Dataset"
    dashboard["A1"].font = Font(bold=True, size=18, color="111827")
    dashboard["A3"] = "Total candidates"
    dashboard["B3"] = summary["count"]
    dashboard["A4"] = "Source"
    dashboard["B4"] = summary["source"]
    dashboard["A5"] = "Product import CSV"
    dashboard["B5"] = summary["import_csv"]
    dashboard["A7"] = "Risk Level"
    dashboard["B7"] = "Count"
    row = 8
    for key, value in summary["risk_level_counts"].items():
        dashboard[f"A{row}"] = key
        dashboard[f"B{row}"] = value
        row += 1
    dashboard["D7"] = "Priority"
    dashboard["E7"] = "Count"
    row = 8
    for key, value in summary["priority_counts"].items():
        dashboard[f"D{row}"] = key
        dashboard[f"E{row}"] = value
        row += 1
    for cell in dashboard["A7:B7"][0] + dashboard["D7:E7"][0]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
    chart = BarChart()
    chart.title = "Candidates By Risk Level"
    chart.y_axis.title = "Count"
    chart.x_axis.title = "Risk Level"
    data = Reference(dashboard, min_col=2, min_row=7, max_row=11)
    cats = Reference(dashboard, min_col=1, min_row=8, max_row=11)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    dashboard.add_chart(chart, "G3")
    dashboard.column_dimensions["A"].width = 28
    dashboard.column_dimensions["B"].width = 32
    dashboard.column_dimensions["D"].width = 18
    dashboard.column_dimensions["E"].width = 12

    import_ws = wb.create_sheet("Import_CSV")
    write_sheet(import_ws, list(import_rows[0].keys()), import_rows)

    full_ws = wb.create_sheet("Full_Dataset")
    write_sheet(full_ws, list(full_rows[0].keys()), full_rows)

    dictionary = wb.create_sheet("Data_Dictionary")
    dictionary_rows = [
        ("name", "Candidate display name; synthetic."),
        ("company_name", "Current company used by TalentNet import."),
        ("industry", "Company industry/category."),
        ("website", "Company website."),
        ("title", "Current role title."),
        ("department", "Functional team."),
        ("location", "Primary work location."),
        ("domain", "Technical domain."),
        ("profile_summary", "Candidate narrative for product display and LLM context."),
        ("source_url", "Synthetic Google search URL, not scraped personal data."),
        ("risk_score", "Synthetic hunting priority score from 0 to 100."),
        ("confidence_score", "Synthetic evidence confidence score from 0 to 100."),
        ("risk_level", "low, medium, high, or very_high."),
        ("hunt_priority", "A/B/C action priority."),
        ("recommended_action", "Suggested TalentNet workflow action."),
    ]
    dictionary.append(["Field", "Meaning"])
    for item in dictionary_rows:
        dictionary.append(list(item))
    for cell in dictionary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    dictionary.column_dimensions["A"].width = 28
    dictionary.column_dimensions["B"].width = 86
    dictionary.freeze_panes = "A2"

    wb.save(OUTPUT_XLSX)
    print(OUTPUT_XLSX)


if __name__ == "__main__":
    main()
